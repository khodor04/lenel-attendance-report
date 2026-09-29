"""
Milestone XProtect — Camera Inventory & Storage Report
Sheets:
  1. Camera Inventory — full details per camera
  2. Summary by Recording Server
"""

import pyodbc
import openpyxl
import xml.etree.ElementTree as ET
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from pathlib import Path

SERVER   = 'localhost'
DATABASE = 'surveillance'
CONN_STR = (f"DRIVER={{ODBC Driver 17 for SQL Server}};"
            f"SERVER={SERVER};DATABASE={DATABASE};"
            f"Trusted_Connection=yes;TrustServerCertificate=yes;")

CODEC_MAP = {'0': 'MJPEG', '1': 'MPEG4', '6': 'H.264', '8': 'H.265', '11': 'H.265+'}

# ── Queries ───────────────────────────────────────────────────────────────────
CAMERA_QUERY = """
SELECT
    d.Name                                                          AS CameraName,
    REPLACE(REPLACE(ISNULL(h.URI,''), 'http://', ''), '/', '')     AS IPAddress,
    ISNULL(ddi.macaddress, '')                                      AS MACAddress,
    rec.Name                                                        AS RecordingServer,
    rs.name                                                         AS StorageName,
    rs.path                                                         AS StoragePath,
    CAST(rs.max_size_mb AS BIGINT)                                  AS MaxSizeMB,
    rs.retain_minutes                                               AS RetainMinutes,
    rsd.recording_enabled                                           AS RecordingEnabled,
    rsd.recording_framerate                                         AS MilestoneFPSCap,
    CAST(d.IDDevice  AS NVARCHAR(36))                               AS DeviceGUID,
    CAST(d.Enabled   AS INT)                                        AS DeviceEnabled,
    CAST(d.Settings  AS NVARCHAR(MAX))                              AS DeviceSettingsXML,
    CAST(d.Streams   AS NVARCHAR(MAX))                              AS StreamsXML,
    CAST(h.Settings  AS NVARCHAR(MAX))                              AS HWSettingsXML
FROM dbo.Devices d
INNER JOIN dbo.Hardware               h   ON d.IDHardware  = h.IDHardware
INNER JOIN dbo.Recorders              rec ON h.IDRecorder  = rec.IDRecorder
INNER JOIN dbo.RecordingStorageDevice rsd ON rsd.device_id = d.IDDevice
INNER JOIN dbo.RecordingStorage       rs  ON rs.pkid       = rsd.recordingstorage_id
LEFT  JOIN dbo.DetectedDeviceInfo     ddi
    ON ddi.ipaddress = REPLACE(REPLACE(ISNULL(h.URI,''), 'http://', ''), '/', '')
WHERE d.DeviceType = 'Camera'
ORDER BY rec.Name, d.Name
"""

STORAGE_QUERY = """
SELECT LOWER(instance_name) AS instance_name, calculated_value AS used_bytes
FROM (
    SELECT pci.instance_name, pcv.calculated_value,
           ROW_NUMBER() OVER (PARTITION BY pci.instance_name ORDER BY pcv.collection_time DESC) AS rn
    FROM dbo.PerformanceCounterInstance pci
    INNER JOIN dbo.PerformanceCounterValue pcv ON pcv.instance_id = pci.pkid
    WHERE pci.category_name = 'VideoOS Recording Server Device Storage'
      AND pci.counter_name  = 'Used Bytes'
) sub WHERE rn = 1
"""

# ── XML helpers ───────────────────────────────────────────────────────────────
def _xml_setting(root, name):
    for s in root.iter('setting'):
        n = s.find('name')
        v = s.find('value')
        if n is not None and n.text == name and v is not None:
            return (v.text or '').strip()
    return ''


def parse_hw_settings(xml_str):
    """Extract model, firmware, serial from Hardware.Settings XML."""
    out = {'model': '', 'firmware': '', 'serial': ''}
    if not xml_str:
        return out
    try:
        root = ET.fromstring(xml_str)
        out['model']    = _xml_setting(root, 'DetectedModelName')
        out['firmware'] = _xml_setting(root, 'FirmwareVersion')
        out['serial']   = _xml_setting(root, 'SerialNumber')
    except ET.ParseError:
        pass
    return out


def parse_device_settings(dev_xml, streams_xml):
    """
    Find the recording stream's hwid from StreamsXML,
    then pull Codec, Resolution, FPS, Bitrate, PTZ from DeviceSettingsXML.
    """
    out = {'codec': '', 'resolution': '', 'cam_fps': '', 'bitrate_kbps': '', 'ptz': ''}
    if not dev_xml:
        return out

    # Find recording stream hwid
    rec_hwid = None
    if streams_xml:
        try:
            sr = ET.fromstring(streams_xml)
            for stream in sr.findall('stream'):
                rec = stream.find('record')
                hwid = stream.find('hwid')
                if rec is not None and rec.text == 'True' and hwid is not None:
                    rec_hwid = hwid.text  # e.g. "stream:0.0.1"
                    break
        except ET.ParseError:
            pass

    try:
        root = ET.fromstring(dev_xml)
        # PTZ
        out['ptz'] = _xml_setting(root, 'PTZEnabled').capitalize() or 'False'

        # Recording stream settings
        for settings_block in root.findall('.//settings'):
            hwid_attr = settings_block.get('hwid', '')
            if rec_hwid and hwid_attr != rec_hwid:
                continue
            if not hwid_attr.startswith('stream:'):
                continue
            codec_raw = _xml_setting(settings_block, 'Codec')
            out['codec']        = CODEC_MAP.get(codec_raw, codec_raw or '')
            out['resolution']   = _xml_setting(settings_block, 'Resolution')
            out['cam_fps']      = _xml_setting(settings_block, 'FPS')
            bitrate             = _xml_setting(settings_block, 'Bitrate')
            out['bitrate_kbps'] = bitrate
            break
    except ET.ParseError:
        pass
    return out


# ── Formatting helpers ────────────────────────────────────────────────────────
def fmt_mac(raw):
    m = (raw or '').strip().upper().replace(':', '').replace('-', '')
    if len(m) == 12:
        return ':'.join(m[i:i+2] for i in range(0, 12, 2))
    return raw or ''


def bytes_to_readable(b):
    if b is None:
        return ''
    b = float(b)
    for unit in ('B', 'KB', 'MB', 'GB', 'TB'):
        if b < 1024:
            return f"{b:.2f} {unit}"
        b /= 1024
    return f"{b:.2f} PB"


def _fill(hex_color):
    return PatternFill('solid', fgColor=hex_color)


def _border():
    s = Side(style='thin', color='CCCCCC')
    return Border(left=s, right=s, top=s, bottom=s)


# ── Data fetch ────────────────────────────────────────────────────────────────
def fetch_data():
    with pyodbc.connect(CONN_STR, timeout=10) as conn:
        cursor = conn.cursor()

        cursor.execute(CAMERA_QUERY)
        cols    = [c[0] for c in cursor.description]
        cameras = [dict(zip(cols, row)) for row in cursor.fetchall()]

        cursor.execute(STORAGE_QUERY)
        storage_map = {}
        for row in cursor.fetchall():
            name = (row.instance_name or '').lower()
            if '[' in name and name.endswith(']'):
                guid = name.split('[')[-1].rstrip(']')
                storage_map[guid] = row.used_bytes

    for cam in cameras:
        cam['UsedBytes'] = storage_map.get(cam['DeviceGUID'].lower())
        cam.update(parse_hw_settings(cam.get('HWSettingsXML')))
        cam.update(parse_device_settings(cam.get('DeviceSettingsXML'), cam.get('StreamsXML')))

    return cameras


# ── Excel builder ─────────────────────────────────────────────────────────────
def build_excel(cameras, out_path):
    wb = openpyxl.Workbook()

    # ── Sheet 1 ───────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Camera Inventory'

    title_font  = Font(name='Segoe UI', size=14, bold=True, color='1F3864')
    hdr_font    = Font(name='Segoe UI', size=9,  bold=True, color='FFFFFF')
    data_font   = Font(name='Segoe UI', size=9)
    hdr_fill    = _fill('1F3864')
    alt_fill    = _fill('EEF2F7')
    bdr         = _border()
    center      = Alignment(horizontal='center', vertical='center')
    left        = Alignment(horizontal='left',   vertical='center')

    ws.merge_cells('A1:T1')
    ws['A1'] = 'Milestone XProtect — Camera Inventory Report'
    ws['A1'].font = title_font
    ws['A1'].alignment = left
    ws.row_dimensions[1].height = 26

    ws.merge_cells('A2:T2')
    ws['A2'] = (f"Generated: {datetime.now().strftime('%d %b %Y  %H:%M')}"
                f"    Total Cameras: {len(cameras)}")
    ws['A2'].font      = Font(name='Segoe UI', size=9, italic=True, color='555555')
    ws['A2'].alignment = left
    ws.row_dimensions[2].height = 16

    headers = [
        '#', 'Camera Name', 'IP Address', 'MAC Address',
        'Model', 'Serial Number', 'Firmware',
        'Recording Server', 'Storage Name', 'Storage Path',
        'Max Size', 'Retention\n(days)', 'Milestone\nFPS Cap',
        'Cam FPS', 'Codec', 'Resolution', 'Bitrate\n(kbps)',
        'PTZ', 'Recording', 'Used Storage',
    ]
    col_widths = [
        5, 38, 16, 20,
        42, 20, 22,
        22, 18, 16,
        12, 12, 12,
        10, 10, 14, 12,
        8, 11, 16,
    ]

    HDR_ROW = 3
    for ci, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=HDR_ROW, column=ci, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border    = bdr
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.row_dimensions[HDR_ROW].height = 28
    ws.freeze_panes = 'A4'

    recorder_totals = {}

    for i, cam in enumerate(cameras, start=1):
        row  = HDR_ROW + i
        fill = alt_fill if i % 2 == 0 else None

        max_gb   = (cam['MaxSizeMB'] or 0) / 1024
        ret_days = round((cam['RetainMinutes'] or 0) / 60 / 24, 1)
        rec_en   = 'Yes' if cam['RecordingEnabled'] else 'No'
        dev_en   = 'Yes' if cam['DeviceEnabled']    else 'No'
        ptz_val  = 'Yes' if str(cam.get('ptz', '')).lower() == 'true' else 'No'

        row_data = [
            i,
            cam['CameraName'],
            cam['IPAddress'],
            fmt_mac(cam['MACAddress']),
            cam.get('model', ''),
            cam.get('serial', ''),
            cam.get('firmware', ''),
            cam['RecordingServer'],
            cam['StorageName'],
            cam['StoragePath'],
            f"{max_gb:,.0f} GB",
            ret_days,
            cam['MilestoneFPSCap'],
            cam.get('cam_fps', ''),
            cam.get('codec', ''),
            cam.get('resolution', ''),
            cam.get('bitrate_kbps', ''),
            ptz_val,
            rec_en,
            bytes_to_readable(cam['UsedBytes']),
        ]

        CENTER_COLS = {1, 11, 12, 13, 14, 16, 17, 18, 19}
        for ci, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=ci, value=val)
            cell.font      = data_font
            cell.border    = bdr
            cell.alignment = center if ci in CENTER_COLS else left
            if fill:
                cell.fill = fill

        # Colour-code Recording (col 19)
        rc = ws.cell(row=row, column=19)
        rc.font = Font(name='Segoe UI', size=9, bold=True,
                       color='27AE60' if cam['RecordingEnabled'] else 'E74C3C')

        # Colour-code PTZ (col 18)
        pc = ws.cell(row=row, column=18)
        pc.font = Font(name='Segoe UI', size=9, bold=True,
                       color='2980B9' if ptz_val == 'Yes' else '888888')

        rec_name = cam['RecordingServer']
        if rec_name not in recorder_totals:
            recorder_totals[rec_name] = {'cameras': 0, 'used_bytes': 0, 'max_mb': 0}
        recorder_totals[rec_name]['cameras']    += 1
        recorder_totals[rec_name]['used_bytes'] += (cam['UsedBytes'] or 0)
        recorder_totals[rec_name]['max_mb']     += (cam['MaxSizeMB'] or 0)

    # ── Sheet 2: Summary ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet('Summary by Recorder')
    ws2.merge_cells('A1:F1')
    ws2['A1'] = 'Summary by Recording Server'
    ws2['A1'].font      = title_font
    ws2['A1'].alignment = left
    ws2.row_dimensions[1].height = 26

    s_hdrs  = ['Recording Server', 'Cameras', 'Used Storage', 'Max Allocated', '% Used']
    s_widths = [30, 12, 20, 20, 12]
    for ci, (h, w) in enumerate(zip(s_hdrs, s_widths), start=1):
        cell = ws2.cell(row=2, column=ci, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = center; cell.border = bdr
        ws2.column_dimensions[get_column_letter(ci)].width = w

    grand = {'cameras': 0, 'used_bytes': 0, 'max_mb': 0}
    for i, (rname, t) in enumerate(sorted(recorder_totals.items()), start=1):
        row    = 2 + i
        fill   = alt_fill if i % 2 == 0 else None
        max_gb = t['max_mb'] / 1024
        pct    = (t['used_bytes'] / (t['max_mb'] * 1024 * 1024) * 100) if t['max_mb'] else 0
        for ci, val in enumerate(
            [rname, t['cameras'], bytes_to_readable(t['used_bytes']),
             f"{max_gb:,.0f} GB", f"{pct:.1f}%"], start=1
        ):
            cell = ws2.cell(row=row, column=ci, value=val)
            cell.font = data_font; cell.border = bdr
            cell.alignment = left if ci == 1 else center
            if fill: cell.fill = fill
        for k in grand: grand[k] += t[k]

    tr = 2 + len(recorder_totals) + 1
    tf = _fill('2E4B8A')
    g_pct = (grand['used_bytes'] / (grand['max_mb'] * 1024 * 1024) * 100) if grand['max_mb'] else 0
    for ci, val in enumerate(
        ['TOTAL', grand['cameras'], bytes_to_readable(grand['used_bytes']),
         f"{grand['max_mb']/1024:,.0f} GB", f"{g_pct:.1f}%"], start=1
    ):
        cell = ws2.cell(row=tr, column=ci, value=val)
        cell.font = Font(name='Segoe UI', size=9, bold=True, color='FFFFFF')
        cell.fill = tf; cell.border = bdr
        cell.alignment = left if ci == 1 else center

    wb.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == '__main__':
    print("Fetching data…")
    cameras = fetch_data()
    print(f"  {len(cameras)} cameras")
    out = Path.home() / 'Documents' / f"Milestone_Camera_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    build_excel(cameras, str(out))
